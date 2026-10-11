from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
source = ROOT / 'Almaty_PM25_Research_Paper_with_Graphs.docx'
output = ROOT / 'Almaty_PM25_Research_Paper_APA.docx'
doc = Document(source)
start = next(i for i, p in enumerate(doc.paragraphs) if p.text == 'References')
end = next(i for i, p in enumerate(doc.paragraphs) if p.text.startswith('Appendix A'))
entries = [p for p in doc.paragraphs[start + 1:end] if p.text.strip()]
assert len(entries) == 4

# APA 7 author/date order, sentence case, italicized standalone works and
# proceedings titles, and DOI URLs. Retain the requested 1.5 line spacing.
references = [
    [('AirData.kz. (n.d.). ', False),
     ('AirDatakz OpenData: Open air quality data for Kazakhstan', True),
     (' [Data set]. Global Shapers Almaty Hub. Retrieved October 10, 2026, from https://github.com/qazybekb/AirDatakz-OpenData', False)],
    [('Chen, T., & Guestrin, C. (2016). XGBoost: A scalable tree boosting system. In ', False),
     ('Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining', True),
     (' (pp. 785–794). Association for Computing Machinery. https://doi.org/10.1145/2939672.2939785', False)],
    [('Hyndman, R. J., & Athanasopoulos, G. (2021). ', False),
     ('Forecasting: Principles and practice', True),
     (' (3rd ed.). OTexts. https://otexts.com/fpp3/', False)],
    [('scikit-learn developers. (n.d.). ', False),
     ('Ridge', True),
     ('. scikit-learn 1.7.2 documentation. https://scikit-learn.org/1.7/modules/generated/sklearn.linear_model.Ridge.html', False)],
]
for paragraph, runs in zip(entries, references):
    paragraph.clear()
    paragraph.style = doc.styles['Normal']
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = False
    for text, italic in runs:
        run = paragraph.add_run(text)
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)

heading = doc.paragraphs[start]
heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
heading.paragraph_format.page_break_before = True
heading.paragraph_format.space_before = Pt(0)
doc.paragraphs[end].paragraph_format.page_break_before = True
doc.save(output)

# Check that the edit preserves the complete body and all six figures.
before = Document(source)
after = Document(output)
assert len(after.inline_shapes) == len(before.inline_shapes) == 6
assert [p.text for p in before.paragraphs[:start + 1]] == [p.text for p in after.paragraphs[:start + 1]]
assert [p.text for p in before.paragraphs[end:]] == [p.text for p in after.paragraphs[end:]]
assert all(p.paragraph_format.line_spacing == 1.5 for p in after.paragraphs[start + 1:end] if p.text.strip())

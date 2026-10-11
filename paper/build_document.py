from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(1)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption'):
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    fonts = style.element.get_or_add_rPr().rFonts
    for attribute in list(fonts.attrib):
        if 'Theme' in attribute:
            del fonts.attrib[attribute]
    for attribute in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
        fonts.set(qn('w:' + attribute), 'Times New Roman')
    properties = style.element.get_or_add_pPr()
    for border in list(properties.findall(qn('w:pBdr'))):
        properties.remove(border)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)
    if name in ('Title', 'Heading 1', 'Heading 2'):
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
doc.core_properties.title = 'Short term PM2 5 forecasting in Almaty using open station data'
doc.core_properties.author = ''
lines = (ROOT / 'Almaty_PM25_Research_Paper.md').read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [part.strip() for part in lines[i].strip('|').split('|')]
            if not all(set(cell) <= set('-: ') for cell in cells):
                rows.append(cells)
            i += 1
        table = doc.add_table(rows=1, cols=len(rows[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = [0.8, 2.0, 0.8, 0.9, 1.0, 1.0] if len(rows[0]) == 6 else [0.8, 1.95, 1.25, 1.25, 1.25]
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for number, values in enumerate(rows):
            row = table.rows[0] if number == 0 else table.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if number == 0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for cell, value, width in zip(row.cells, values, widths):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = value
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.5
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.space_before = Pt(3)
                    for run in paragraph.runs:
                        run.font.name, run.font.size = 'Times New Roman', Pt(12)
                        run.bold = number == 0
                if number == 0:
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), 'EEEEEE')
                    cell._tc.get_or_add_tcPr().append(shading)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 1')
    elif line.startswith('### '):
        doc.add_paragraph(line[4:], 'Heading 2')
    elif re.match(r'Table \d+\.', line):
        paragraph = doc.add_paragraph(line, 'Caption')
        paragraph.paragraph_format.keep_with_next = True
    else:
        doc.add_paragraph(line)
    i += 1
footer = section.footer.paragraphs[0]
footer.paragraph_format.line_spacing = 1.5
footer.alignment = 2
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for paragraph in list(doc.paragraphs) + [footer]:
    for run in paragraph.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        fonts = run._element.get_or_add_rPr().rFonts
        for attribute in list(fonts.attrib):
            if 'Theme' in attribute:
                del fonts.attrib[attribute]
    properties = paragraph._p.get_or_add_pPr()
    for border in list(properties.findall(qn('w:pBdr'))):
        properties.remove(border)
doc.save(ROOT / 'Almaty_PM25_Research_Paper.docx')

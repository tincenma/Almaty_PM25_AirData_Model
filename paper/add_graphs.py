"""Insert figures into a copy while preserving the existing document text."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent
doc = Document(ROOT / 'Almaty_PM25_Research_Paper.docx')

def paragraph_start(prefix):
    return next(p._p for p in doc.paragraphs if p.text.startswith(prefix))

def insert(anchor, name, caption, alt):
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    picture.paragraph_format.line_spacing = 1.5
    picture.paragraph_format.space_after = Pt(3)
    run = picture.add_run()
    shape = run.add_picture(str(ROOT / 'figures' / name), width=Inches(6.3))
    shape._inline.docPr.set('descr', alt)
    explanation = doc.add_paragraph(caption, 'Caption')
    explanation.paragraph_format.keep_with_next = False
    explanation.paragraph_format.keep_together = True
    explanation.paragraph_format.line_spacing = 1.5
    explanation.paragraph_format.space_after = Pt(8)
    for run in explanation.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.italic = False
    anchor.addnext(picture._p)
    picture._p.addnext(explanation._p)

insert(paragraph_start('The updated experiment uses'), '01_data_coverage.png',
       'Figure 1. PCP #31 monthly observation coverage. The first and last months are partial. Missing hours remain missing. Source: saved monthly_data_coverage.csv.',
       'Monthly coverage decreases across parts of the PCP 31 snapshot; the grid contains 65.6 percent observed hours overall.')
insert(doc.tables[1]._tbl, '02_overall_mae.png',
       'Figure 2. Held-out MAE by horizon for each station. Lower values indicate smaller average errors. Different stations and periods limit comparisons between panels. Source: saved test_metrics.csv files.',
       'XGBoost has lower MAE than persistence at each horizon in both separate experiments, with the greatest relative improvement at six hours.')
insert(paragraph_start('The six-hour paired MAE gain'), '03_uncertainty.png',
       'Figure 3. PCP #31 paired MAE gains with approximate 95% weekly-block intervals. Positive values favour XGBoost. Only the six-hour interval excludes zero. Source: paired_block_bootstrap.csv.',
       'One-hour and 24-hour improvement intervals include zero; the six-hour interval ranges from 1.41 to 4.30 micrograms per cubic metre.')
insert(paragraph_start('PCP #31 has no eligible summer'), '04_season_errors.png',
       'Figure 4. PCP #31 next-day MAE by target season with sample counts. Summer has no eligible examples. XGBoost loses to persistence in spring. Source: season_errors.csv.',
       'Winter next-day errors exceed autumn and spring errors, and spring persistence outperforms XGBoost.')
insert(paragraph_start('Sparse weekly history produces'), '05_coverage_errors.png',
       'Figure 5. PCP #31 next-day MAE by observed weekly coverage with sample counts. Bins contain at most 42, 43–126, and 127–168 observed hours. Source: coverage_errors.csv.',
       'XGBoost next-day MAE exceeds persistence in both sparse weekly coverage bins, despite the daily eligibility requirement.')
insert(paragraph_start('For 360 historical next-day peak'), '06_peak_errors.png',
       'Figure 6. Next-day MAE below and above each station’s pre-test 90th percentile. Panels use different vertical scales and cutoffs. These are statistical thresholds. Sources: saved peak_errors.csv files.',
       'On US Embassy peak targets XGBoost is worse than persistence. On PCP 31 peak targets XGBoost is better than persistence but errors remain large.')

doc.save(ROOT / 'Almaty_PM25_Research_Paper_with_Graphs.docx')
